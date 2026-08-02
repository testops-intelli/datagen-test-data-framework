"""
export.py

Exports DataGen mode output (rows fetched via db.execute_and_fetch_all)
to a .xlsx workbook the user can open, filter, and hand off as test
data / validation evidence.
"""

from pathlib import Path
from typing import Any, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font

from .config import EXPORT_DIR


def _autosize_columns(ws, columns: Sequence[str], rows: Sequence[Tuple[Any, ...]]) -> None:
    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(str(col_name))
        for row in rows:
            value = row[col_idx - 1]
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def export_rows_to_xlsx(
    columns: Sequence[str],
    rows: Sequence[Tuple[Any, ...]],
    file_name: str,
    sheet_title: str = "DataGen Output",
) -> Path:
    """
    Writes columns/rows (as returned by db.execute_and_fetch_all) to
    an .xlsx file under EXPORT_DIR. Returns the path written.
    """
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel sheet name limit

    ws.append(list(columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        # jsonb / dict-like cell values aren't natively writable;
        # stringify anything that isn't a plain scalar.
        safe_row = [
            v if isinstance(v, (str, int, float, bool)) or v is None else str(v)
            for v in row
        ]
        ws.append(safe_row)

    ws.freeze_panes = "A2"
    _autosize_columns(ws, columns, rows)

    out_path = EXPORT_DIR / file_name
    wb.save(out_path)
    return out_path
